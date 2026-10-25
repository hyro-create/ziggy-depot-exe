"""Vérifie que toutes les dépendances et ressources nécessaires sont présentes.
Utilisation :  python selfcheck.py     (avant de compiler)
               Ziggy-diagnostic.exe --check   (sur l'application compilée)"""
import importlib, sys
from pathlib import Path
from paths import DATA_DIR, FROZEN, resource_path

REQUIRED = {  # module importé -> paquet pip
    "flask": "flask", "waitress": "waitress", "dotenv": "python-dotenv", "requests": "requests", "httpx": "httpx",
    "google.genai": "google-genai", "google.oauth2.credentials": "google-auth",
    "google_auth_oauthlib.flow": "google-auth-oauthlib", "googleapiclient.discovery": "google-api-python-client",
    "msal": "msal", "spotipy": "spotipy", "pyperclip": "pyperclip", "PIL": "Pillow",
    "docx": "python-docx", "openpyxl": "openpyxl", "pptx": "python-pptx", "pypdf": "pypdf",
}
OWN = ["version", "paths", "logsetup", "secrets_store", "settings", "database", "gemini", "security", "actions",
       "reminders", "system_integration", "readers", "extras", "apps", "telegram_bot", "google_services", "microsoft_services"]
RESOURCES = ("static/index.html", "static/aide.html", "static/logo.png", "static/logo-light.png", "static/favicon.ico",
             "assets/ziggy_icon.png", ".env.example")


def check() -> list:
    problems = []
    for mod, pkg in REQUIRED.items():
        try:
            importlib.import_module(mod)
        except Exception as e:
            problems.append(f"Dépendance manquante ou cassée : {pkg}  (import {mod} : {type(e).__name__}: {e})")
    if sys.platform == "win32":                     # la barre des tâches n'existe que sous Windows
        try:
            import pystray  # noqa
        except Exception as e:
            problems.append(f"Dépendance manquante : pystray ({type(e).__name__}: {e})")
    for mod in OWN:
        try:
            importlib.import_module(mod)
        except Exception as e:
            problems.append(f"Module Ziggy illisible : {mod}.py ({type(e).__name__}: {e})")
    for rel in RESOURCES:
        if not resource_path(rel).is_file():
            problems.append(f"Ressource absente de l'application : {rel}")
    try:
        import certifi
        if not Path(certifi.where()).is_file():
            problems.append("Certificats HTTPS (certifi) absents : les connexions Internet échoueront.")
    except Exception as e:
        problems.append(f"certifi : {e}")
    try:
        from googleapiclient import discovery_cache
        for api in ("gmail.v1", "calendar.v3", "drive.v3"):
            if not (Path(discovery_cache.__file__).parent / "documents" / f"{api}.json").is_file():
                problems.append(f"Description d'API Google absente : {api} (Gmail/Agenda/Drive ne marcheront pas).")
    except Exception as e:
        problems.append(f"googleapiclient : {e}")
    try:
        import docx, pptx
        docx.Document(); pptx.Presentation()       # échoue si les modèles .docx/.pptx intégrés manquent
    except Exception as e:
        problems.append(f"Modèles Word/PowerPoint absents (lecture des documents impossible) : {e}")
    try:
        import io, openpyxl
        openpyxl.Workbook().save(io.BytesIO())
    except Exception as e:
        problems.append(f"Excel : {e}")
    try:                                            # le SDK Gemini doit exposer ce que Ziggy utilise
        from google.genai import errors, types
        assert hasattr(errors, "APIError") and "supported_actions" in types.Model.model_fields
    except Exception as e:
        problems.append(f"Version incompatible du SDK google-genai : {type(e).__name__} {e}")
    t = DATA_DIR / ".ziggy_check"
    try:
        t.write_text("x"); t.unlink()
    except OSError as e:
        problems.append(f"Dossier de données non inscriptible : {DATA_DIR} ({e})")
    return problems


def run() -> int:
    print(f"Mode : {'application compilée' if FROZEN else 'script Python'} — Python {sys.version.split()[0]}")
    print(f"Données (mémoire, paramètres, clé chiffrée, journaux) : {DATA_DIR}")
    p = check()
    if p:
        print(f"\n✖ {len(p)} problème(s) :")
        for x in p:
            print("  -", x)
        return 1
    print("✔ Toutes les dépendances et ressources sont présentes.")
    return 0


if __name__ == "__main__":
    sys.exit(run())
