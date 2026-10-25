"""Lecture de texte dans Word, Excel, PowerPoint, PDF et fichiers texte."""
import io

LIMIT = 20000


def extract_text(name: str, data: bytes) -> str:
    n = name.lower()
    try:
        if n.endswith(".docx"):
            from docx import Document
            d = Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            text = "\n".join(parts)
        elif n.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(f"## Feuille : {ws.title}")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 300:
                        parts.append("… (tronqué)"); break
                    if any(c is not None for c in row):
                        parts.append(" | ".join("" if c is None else str(c) for c in row))
            text = "\n".join(parts)
        elif n.endswith(".pptx"):
            from pptx import Presentation
            parts = []
            for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
                parts.append(f"## Diapo {i}")
                for sh in slide.shapes:
                    if sh.has_text_frame and sh.text_frame.text.strip():
                        parts.append(sh.text_frame.text)
            text = "\n".join(parts)
        elif n.endswith(".pdf"):
            from pypdf import PdfReader
            text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
            if not text.strip():
                return "PDF sans texte (scan/image) : lecture impossible."
        elif n.endswith((".doc", ".xls", ".ppt")):
            return "Ancien format Office (.doc/.xls/.ppt) non supporté : demande à l'utilisateur de l'enregistrer en .docx/.xlsx/.pptx."
        else:
            text = data.decode("utf-8", errors="replace")
        return text[:LIMIT]
    except Exception as e:
        return f"Erreur de lecture ({name}) : {e}"
